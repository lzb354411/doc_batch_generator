# -*- coding: utf-8 -*-
"""金标准对比：逐文件对比 Go 版输出（testdata/out）与 Python 版输出（testdata/out_python）。

对比内容：
  1. 相对路径集合完全一致
  2. txt/csv：字节级一致
  3. xlsx：逐工作表逐单元格值一致（含数字、日期格式化结果）
  4. docx：正文段落 + 表格段落文本一致

用法（在 e:\\workgo\\work\\doc_batch_generator 下执行）：
    python docgen/cmd/pygolden/compare.py
"""
import os
import sys

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
GO_OUT = os.path.join(BASE, "testdata", "out")
PY_OUT = os.path.join(BASE, "testdata", "out_python")

failures = []
checked = 0


def fail(rel, reason):
    failures.append("%s：%s" % (rel, reason))


def rel_files(root):
    result = set()
    for dirpath, _dirnames, filenames in os.walk(root):
        for name in filenames:
            result.add(os.path.relpath(os.path.join(dirpath, name), root))
    return result


def load_xlsx_cells(path):
    """返回 {sheet: [(cell_ref, value_str)]}，按行序展开"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    data = {}
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    v = c.value
                    if hasattr(v, "strftime"):  # datetime/date/time
                        v = v.strftime("%Y-%m-%d %H:%M:%S") if hasattr(v, "hour") else v.strftime("%Y-%m-%d")
                    cells.append((c.coordinate, str(v)))
        data[ws.title] = cells
    wb.close()
    return data


def load_docx_texts(path):
    """返回正文段落文本列表 + 表格内段落文本列表（用 python-docx）"""
    import docx
    d = docx.Document(path)
    texts = [("body", i, p.text) for i, p in enumerate(d.paragraphs)]
    for ti, table in enumerate(d.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                for pi, p in enumerate(cell.paragraphs):
                    texts.append(("tbl%d-r%d-c%d-p%d" % (ti, ri, ci, pi), -1, p.text))
    return texts


def main():
    global checked
    go_files = rel_files(GO_OUT)
    py_files = rel_files(PY_OUT)

    if go_files != py_files:
        only_go = sorted(go_files - py_files)
        only_py = sorted(py_files - go_files)
        if only_go:
            fail("(文件集合)", "仅 Go 版有：%s" % "; ".join(only_go))
        if only_py:
            fail("(文件集合)", "仅 Python 版有：%s" % "; ".join(only_py))
    checked += 1

    for rel in sorted(go_files & py_files):
        gp = os.path.join(GO_OUT, rel)
        pp = os.path.join(PY_OUT, rel)
        ext = os.path.splitext(rel)[1].lower()

        if ext in (".txt", ".csv"):
            with open(gp, "rb") as f:
                gb = f.read()
            with open(pp, "rb") as f:
                pb = f.read()
            if gb != pb:
                fail(rel, "字节不一致\n    Go     : %r\n    Python : %r" % (gb[:300], pb[:300]))
            checked += 1

        elif ext in (".xlsx", ".xlsm"):
            gc, pc = load_xlsx_cells(gp), load_xlsx_cells(pp)
            if set(gc) != set(pc):
                fail(rel, "工作表集合不一致：Go=%s Python=%s" % (sorted(gc), sorted(pc)))
            else:
                for sheet in gc:
                    if gc[sheet] != pc[sheet]:
                        diff = [("Go=%r Python=%r" % (a, b))
                                for a, b in zip(gc[sheet], pc[sheet]) if a != b]
                        if len(gc[sheet]) != len(pc[sheet]):
                            diff.append("单元格数量 Go=%d Python=%d" % (len(gc[sheet]), len(pc[sheet])))
                        fail(rel, "工作表[%s]不一致：%s" % (sheet, "; ".join(diff[:5])))
            checked += 1

        elif ext == ".docx":
            gt, pt = load_docx_texts(gp), load_docx_texts(pp)
            if len(gt) != len(pt):
                fail(rel, "段落数不一致：Go=%d Python=%d" % (len(gt), len(pt)))
            else:
                for (gk, _, gv), (pk, _, pv) in zip(gt, pt):
                    if gk != pk or gv != pv:
                        fail(rel, "[%s] 不一致\n    Go     : %r\n    Python : %r" % (gk, gv, pv))
            checked += 1

    print("对比完成：%d 项（1 项文件集合 + %d 个文件）" % (checked, len(go_files & py_files)))
    if failures:
        print("\n==== 发现 %d 处不一致 ====" % len(failures))
        for msg in failures:
            print("[不一致]", msg)
        sys.exit(1)
    print("Go 版与 Python 版输出完全一致")


if __name__ == "__main__":
    main()
